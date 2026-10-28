"""Fayllarni Word (.docx) ga aylantiruvchi yadro.

Qo'llab-quvvatlanadi: PDF, XLSX/XLSM, XLS, CSV/TSV, TXT/MD.
Og'ir ish `asyncio.to_thread` orqali bajariladi — bot qotmaydi.
"""
from __future__ import annotations

import asyncio
import csv
import io
import logging
import re
from datetime import date, datetime
from pathlib import Path

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".xlsx", ".xlsm", ".xls", ".csv", ".tsv", ".txt", ".md"}

MAX_ROWS = 3000
MAX_COLS = 60
MAX_TEXT_LINES = 20000
MAX_PDF_PAGES = 100  # kattaroq PDF matn rejimida (xotirani tejash uchun)


class ConversionError(Exception):
    """Foydalanuvchiga ko'rsatish mumkin bo'lgan xatolik."""


async def convert_file(source: Path | str, target: Path | str) -> Path:
    """Faylni .docx ga o'giradi va natija yo'lini qaytaradi."""
    return await asyncio.to_thread(_convert_sync, Path(source), Path(target))


def _convert_sync(source: Path, target: Path) -> Path:
    ext = source.suffix.lower()
    try:
        if ext == ".pdf":
            _pdf_to_docx(source, target)
        elif ext in (".xlsx", ".xlsm"):
            _excel_to_docx(source, target)
        elif ext == ".xls":
            _xls_to_docx(source, target)
        elif ext in (".csv", ".tsv"):
            _csv_to_docx(source, target)
        elif ext in (".txt", ".md"):
            _text_to_docx(source, target)
        else:
            raise ConversionError("Bu format qo'llab-quvvatlanmaydi.")
    except ConversionError:
        raise
    except Exception as exc:
        logger.exception("Konvertatsiya xatosi: %s", source)
        raise ConversionError(
            "Faylni o'qishda xatolik yuz berdi. "
            "Fayl buzilgan yoki noto'g'ri formatda bo'lishi mumkin."
        ) from exc

    if not target.exists() or target.stat().st_size < 1000:
        raise ConversionError("Natija bo'sh chiqdi. Faylni tekshirib, qayta yuboring.")
    return target


def _new_document(landscape: bool = False):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm, Pt

    document = Document()
    for section in document.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)
        if landscape:
            width, height = section.page_width, section.page_height
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = height, width

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    return document


def _add_heading(document, text: str) -> None:
    text = (text or "").strip()
    if text:
        document.add_heading(text[:200], level=1)


def _format_cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:.10g}"
    return str(value).replace("\x00", "").strip()


def _add_table(document, rows: list[list[str]]) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]

    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.autofit = True

    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(value)[:5000])
            run.font.size = Pt(8)
            if row_index == 0:
                run.bold = True
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "D9E2F3")
                cell._tc.get_or_add_tcPr().append(shade)

    # Jadval sarlavhasi har bir sahifada takrorlansin.
    if len(rows) > 1:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def _normalize_rows(rows: list[list[str]]) -> list[list[str]]:
    clean = []
    for row in rows:
        while row and row[-1] == "":
            row.pop()
        if any(row):
            clean.append(row)
    return clean


def _pdf_to_docx(source: Path, target: Path) -> None:
    page_count = _pdf_page_count(source)
    if page_count and page_count > MAX_PDF_PAGES:
        logger.info("PDF %d sahifa — matn rejimida konvertatsiya", page_count)
        _pdf_text_fallback(source, target)
        return

    try:
        _pdf_with_pdf2docx(source, target)
        if target.exists() and target.stat().st_size > 5000:
            return
        logger.warning("pdf2docx bo'sh natija berdi — matn usuliga o'tamiz")
    except Exception as exc:
        logger.warning("pdf2docx ishlamadi (%s) — matn usuliga o'tamiz", exc)

    _pdf_text_fallback(source, target)


def _pdf_page_count(source: Path) -> int:
    try:
        from pypdf import PdfReader

        return len(PdfReader(str(source)).pages)
    except Exception:
        return 0


def _pdf_with_pdf2docx(source: Path, target: Path) -> None:
    from pdf2docx import Converter

    converter = Converter(str(source))
    try:
        converter.convert(str(target))
    finally:
        converter.close()


def _pdf_text_fallback(source: Path, target: Path) -> None:
    from docx.shared import Pt
    from pypdf import PdfReader

    reader = PdfReader(str(source))
    if reader.is_encrypted:
        try:
            unlocked = reader.decrypt("")
        except Exception as exc:
            raise ConversionError("PDF parol bilan himoyalangan.") from exc
        if not unlocked:
            raise ConversionError("PDF parol bilan himoyalangan. Parolni olib tashlab yuboring.")

    document = _new_document()
    found = False
    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            continue
        found = True
        heading = document.add_heading(f"{page_number}-sahifa", level=3)
        for run in heading.runs:
            run.font.size = Pt(10)
        for line in text.splitlines():
            if line.strip():
                document.add_paragraph(line.rstrip())

    if not found:
        raise ConversionError(
            "PDF ichidan matn topilmadi — fayl skanerlangan bo'lishi mumkin. "
            "Hozircha OCR qo'llab-quvvatlanmaydi."
        )
    document.save(str(target))


def _excel_to_docx(source: Path, target: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(filename=str(source), data_only=True, read_only=True)
    document = _new_document(landscape=True)
    used_sheets = 0
    try:
        for sheet in workbook.worksheets:
            rows: list[list[str]] = []
            truncated = False
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= MAX_ROWS:
                    truncated = True
                    break
                rows.append([_format_cell(value) for value in row[:MAX_COLS]])

            rows = _normalize_rows(rows)
            if not rows:
                continue
            _add_heading(document, sheet.title + (" (qismi)" if truncated else ""))
            _add_table(document, rows)
            document.add_paragraph()
            used_sheets += 1
    finally:
        workbook.close()

    if not used_sheets:
        raise ConversionError("Excel fayl bo'sh — unda ma'lumot topilmadi.")
    document.save(str(target))


def _xls_to_docx(source: Path, target: Path) -> None:
    import xlrd

    workbook = xlrd.open_workbook(str(source))
    document = _new_document(landscape=True)
    used_sheets = 0

    for sheet in workbook.sheets():
        row_count = min(sheet.nrows, MAX_ROWS)
        col_count = min(sheet.ncols, MAX_COLS)
        truncated = sheet.nrows > MAX_ROWS or sheet.ncols > MAX_COLS
        rows: list[list[str]] = []

        for row_index in range(row_count):
            values: list[str] = []
            for col_index in range(col_count):
                cell = sheet.cell(row_index, col_index)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        value = xlrd.xldate_as_datetime(cell.value, workbook.datemode)
                        values.append(_format_cell(value))
                    except Exception:
                        values.append(_format_cell(cell.value))
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    values.append("TRUE" if cell.value else "FALSE")
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                    values.append("")
                else:
                    values.append(_format_cell(cell.value))
            rows.append(values)

        rows = _normalize_rows(rows)
        if not rows:
            continue
        _add_heading(document, sheet.name + (" (qismi)" if truncated else ""))
        _add_table(document, rows)
        document.add_paragraph()
        used_sheets += 1

    if not used_sheets:
        raise ConversionError("Excel fayl bo'sh — unda ma'lumot topilmadi.")
    document.save(str(target))


def _csv_to_docx(source: Path, target: Path) -> None:
    text = _read_text(source)
    if not text.strip():
        raise ConversionError("Fayl bo'sh.")

    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if source.suffix.lower() == ".tsv" else csv.excel

    rows: list[list[str]] = []
    truncated = False
    for index, row in enumerate(csv.reader(io.StringIO(text), dialect)):
        if index >= MAX_ROWS:
            truncated = True
            break
        rows.append([_format_cell(value) for value in row[:MAX_COLS]])

    rows = _normalize_rows(rows)
    if not rows:
        raise ConversionError("Fayl bo'sh — unda ma'lumot topilmadi.")

    document = _new_document(landscape=True)
    _add_heading(document, source.stem + (" (qismi)" if truncated else ""))
    _add_table(document, rows)
    document.save(str(target))


def _read_text(source: Path) -> str:
    data = source.read_bytes()
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    for encoding in ("utf-8-sig", "cp1251", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _text_to_docx(source: Path, target: Path) -> None:
    text = _read_text(source)
    if not text.strip():
        raise ConversionError("Fayl bo'sh.")

    lines = text.splitlines()
    truncated = len(lines) > MAX_TEXT_LINES
    lines = lines[:MAX_TEXT_LINES]

    document = _new_document()
    if source.suffix.lower() == ".md":
        _add_markdown(document, lines)
    else:
        for line in lines:
            document.add_paragraph(line)

    if truncated:
        document.add_paragraph("… (fayl qisqartirildi)")
    document.save(str(target))


_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET_RE = re.compile(r"^\s*[-*+]\s+(.*)$")
_NUMBER_RE = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_INLINE_RE = re.compile(r"(\*\*.+?\*\*|__.+?__|`.+?`|\*.+?\*)")


def _add_markdown(document, lines: list[str]) -> None:
    from docx.shared import Cm

    in_code = False
    code_lines: list[str] = []
    for line in lines:
        if line.strip().startswith("```"):
            if in_code:
                _add_code_block(document, code_lines)
                code_lines = []
            in_code = not in_code
            continue
        if in_code:
            code_lines.append(line)
            continue

        match = _HEADING_RE.match(line)
        if match:
            document.add_heading(match.group(2).strip(), level=min(len(match.group(1)), 4))
            continue

        match = _BULLET_RE.match(line)
        if match:
            paragraph = document.add_paragraph(style="List Bullet")
            _add_inline(paragraph, match.group(1))
            continue

        match = _NUMBER_RE.match(line)
        if match:
            paragraph = document.add_paragraph(style="List Number")
            _add_inline(paragraph, match.group(1))
            continue

        if line.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(1)
            _add_inline(paragraph, line[2:], italic=True)
            continue

        paragraph = document.add_paragraph()
        _add_inline(paragraph, line)

    if in_code and code_lines:
        _add_code_block(document, code_lines)


def _add_code_block(document, lines: list[str]) -> None:
    from docx.shared import Pt

    paragraph = document.add_paragraph()
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def _add_inline(paragraph, text: str, italic: bool = False) -> None:
    position = 0
    for match in _INLINE_RE.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            run.italic = italic
        token = match.group(0)
        if token.startswith("**") or token.startswith("__"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        run.italic = italic
