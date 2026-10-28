"""Konvertorlarning o'z-o'zini tekshiruvi.

Ishga tushirish: python -m unittest discover -s tests -v
"""
from __future__ import annotations

import asyncio
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from converters import (
    ConversionError,
    _pdf_text_fallback,
    convert_file,
)


def _make_pdf(path: Path, text: str) -> None:
    """Tashqi kutubxonasiz bir sahifali, matnli PDF yasaydi."""
    stream = f"BT /F1 18 Tf 72 760 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n"
        + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"

    xref_offset = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n"
    ).encode()
    path.write_bytes(out)


def _document_text(path: Path) -> str:
    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


class ConverterTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def _convert(self, name: str) -> Path:
        source = self.root / name
        target = self.root / (source.stem + ".docx")
        asyncio.run(convert_file(source, target))
        self.assertTrue(target.exists())
        return target

    def test_pdf_to_docx(self) -> None:
        source = self.root / "test.pdf"
        _make_pdf(source, "WordBot PDF test")
        target = self._convert("test.pdf")
        self.assertIn("WordBot", _document_text(target))

    def test_pdf_text_fallback(self) -> None:
        source = self.root / "fallback.pdf"
        target = self.root / "fallback.docx"
        _make_pdf(source, "Fallback matni")
        _pdf_text_fallback(source, target)
        self.assertIn("Fallback matni", _document_text(target))

    def test_xlsx_to_docx(self) -> None:
        source = self.root / "hisobot.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Hisobot"
        sheet.append(["Mahsulot", "Narx", "Soni"])
        sheet.append(["Olma", 12000, 5])
        sheet.append(["Uzum", 25000.5, 2])
        workbook.save(source)

        target = self._convert("hisobot.xlsx")
        text = _document_text(target)
        self.assertIn("Hisobot", text)
        self.assertIn("Olma", text)
        self.assertIn("12000", text)
        self.assertIn("25000.5", text)
        self.assertGreaterEqual(len(Document(str(target)).tables), 1)

    def test_csv_to_docx(self) -> None:
        source = self.root / "odamlar.csv"
        source.write_text("Ism;Yosh\nAli;25\nVali;30\n", encoding="utf-8")
        target = self._convert("odamlar.csv")
        text = _document_text(target)
        self.assertIn("Ali", text)
        self.assertIn("Vali", text)
        self.assertEqual(len(Document(str(target)).tables), 1)

    def test_txt_to_docx(self) -> None:
        source = self.root / "matn.txt"
        source.write_text("Salom dunyo\nIkkinchi qator\n", encoding="utf-8")
        target = self._convert("matn.txt")
        text = _document_text(target)
        self.assertIn("Salom dunyo", text)
        self.assertIn("Ikkinchi qator", text)

    def test_markdown_to_docx(self) -> None:
        source = self.root / "readme.md"
        source.write_text(
            "# Sarlavha\n\n## Kichik sarlavha\n\n- Birinchi\n- **Muhim** band\n",
            encoding="utf-8",
        )
        target = self._convert("readme.md")
        text = _document_text(target)
        self.assertIn("Sarlavha", text)
        self.assertIn("Muhim", text)

    def test_unsupported_format(self) -> None:
        source = self.root / "eski.doc"
        source.write_bytes(b"not a real doc")
        with self.assertRaises(ConversionError):
            asyncio.run(convert_file(source, self.root / "eski.docx"))

    def test_empty_file(self) -> None:
        source = self.root / "bosh.txt"
        source.write_text("   \n", encoding="utf-8")
        with self.assertRaises(ConversionError):
            asyncio.run(convert_file(source, self.root / "bosh.docx"))


if __name__ == "__main__":
    unittest.main()
